"""Self-made .xlsx files for the round-trip loss measurement (docs/xlsx-roundtrip-losses.md).

One file per feature group, so a loss points at its group. Written by openpyxl 3.1 (not by Excel):
what Excel itself writes (pivots, threaded comments, sparklines, slicers, external links) is not
here and needs real documents. Run: /usr/bin/python3 scripts/xlsx-roundtrip/fixtures.py <out dir>
"""
import datetime
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)
FIXED = datetime.datetime(2026, 10, 15, 9, 0, 0)


def save(wb, name):
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.properties.creator = "Sketch Docs fixtures"
    wb.properties.title = name
    wb.save(out / f"{name}.xlsx")


def values():
    wb = Workbook()
    ws = wb.active
    ws.title = "값"
    rows = [["문자", "숫자", "논리", "날짜", "시각", "백분율", "통화"],
            ["사과", 1200, True, datetime.date(2026, 10, 15), datetime.time(9, 30), 0.125, 1234567.891],
            ["배 🍐", -3.5, False, datetime.datetime(1999, 12, 31, 23, 59), datetime.time(0, 0, 1), 1, 0]]
    for row in rows:
        ws.append(row)
    for cell in ws[2] + ws[3]:
        if cell.column == 6:
            cell.number_format = "0.0%"
        if cell.column == 7:
            cell.number_format = '#,##0.00 "원"'
    ws["A5"] = "=1/0"
    ws["A6"] = "긴 글 " * 50
    save(wb, "01-values")
    wb1904 = Workbook()
    wb1904.epoch = datetime.datetime(1904, 1, 1)
    wb1904.active["A1"] = datetime.date(2026, 10, 15)
    save(wb1904, "01b-date1904")


def formulas():
    wb = Workbook()
    ws = wb.active
    ws.title = "수식"
    for i in range(1, 6):
        ws.cell(i, 1, i * 10)
        ws.cell(i, 2, f"=A{i}*2")
    ws["C1"] = "=SUM(A1:A5)"
    ws["C2"] = "=AVERAGE(B1:B5)"
    ws["C3"] = '=IF(C1>100,"크다","작다")'
    ws["C4"] = "=VLOOKUP(30,A1:B5,2,FALSE)"
    ws["C5"] = "=단가*2"
    ws["D1"] = ArrayFormula("D1:D5", "=A1:A5*10")
    other = wb.create_sheet("참조")
    other["A1"] = "=수식!C1+1"
    other["A2"] = "=TODAY()"
    wb.defined_names["단가"] = DefinedName("단가", attr_text="수식!$A$1")
    wb.defined_names["범위"] = DefinedName("범위", attr_text="수식!$A$1:$B$5")
    save(wb, "02-formulas")


def styles():
    wb = Workbook()
    ws = wb.active
    ws.title = "서식"
    thin = Side(style="thin", color="FF336699")
    ws["A1"] = "굵게 빨강"
    ws["A1"].font = Font(name="맑은 고딕", size=14, bold=True, color="FFFF0000")
    ws["A2"] = "기울임 밑줄"
    ws["A2"].font = Font(italic=True, underline="double", strike=True)
    ws["A3"] = "채우기 테두리"
    ws["A3"].fill = PatternFill("solid", fgColor="FFFFF2CC")
    ws["A3"].border = Border(left=thin, right=thin, top=Side(style="double"), bottom=Side(style="dashed", color="FF00AA00"))
    ws["A4"] = "줄 바꿈 가운데 회전"
    ws["A4"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True, text_rotation=45, indent=1)
    ws["A5"] = 0.5
    ws["A5"].number_format = "0.000"
    ws["A6"] = "잠금 해제 숨김"
    ws["A6"].protection = Protection(locked=False, hidden=True)
    ws["A7"] = CellRichText(["보통 ", TextBlock(InlineFont(b=True, color="FF0000FF"), "굵은 파랑"), " 끝"])
    ws["A8"] = "위첨자"
    ws["A8"].font = Font(vertAlign="superscript", name="Arial")
    save(wb, "03-styles")


def layout():
    wb = Workbook()
    ws = wb.active
    ws.title = "배치"
    for r in range(1, 21):
        for c in range(1, 8):
            ws.cell(r, c, r * c)
    ws.merge_cells("A1:C1")
    ws.merge_cells("E5:F8")
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["D"].hidden = True
    ws.row_dimensions[3].height = 40
    ws.row_dimensions[5].hidden = True
    ws.row_dimensions.group(10, 12, outline_level=1, hidden=True)
    ws.column_dimensions.group("F", "G", outline_level=1)
    ws.freeze_panes = "B2"
    ws.sheet_view.zoomScale = 85
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "FF00B050"
    hidden = wb.create_sheet("숨김")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "숨긴 시트"
    very = wb.create_sheet("매우숨김")
    very.sheet_state = "veryHidden"
    rtl = wb.create_sheet("오른쪽부터")
    rtl.sheet_view.rightToLeft = True
    rtl["A1"] = 1
    save(wb, "04-layout")


def rules():
    wb = Workbook()
    ws = wb.active
    ws.title = "규칙"
    for r in range(1, 11):
        ws.cell(r, 1, r * 7 % 10)
        ws.cell(r, 2, r)
    ws.conditional_formatting.add("A1:A10", CellIsRule(operator="greaterThan", formula=["5"], fill=PatternFill("solid", bgColor="FFFFC7CE"), font=Font(color="FF9C0006")))
    ws.conditional_formatting.add("B1:B10", FormulaRule(formula=["MOD(B1,2)=0"], font=Font(bold=True)))
    ws.conditional_formatting.add("C1:C10", ColorScaleRule(start_type="min", start_color="FFF8696B", end_type="max", end_color="FF63BE7B"))
    ws.conditional_formatting.add("D1:D10", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
    ws.conditional_formatting.add("E1:E10", IconSetRule("3TrafficLights1", "percent", [0, 33, 67]))
    for r in range(1, 11):
        for c in (3, 4, 5):
            ws.cell(r, c, r)
    choice = DataValidation(type="list", formula1='"사과,배,감"', allow_blank=True, showErrorMessage=True, errorTitle="오류", error="목록에서 고르세요")
    whole = DataValidation(type="whole", operator="between", formula1="1", formula2="100", promptTitle="입력", prompt="1~100")
    custom = DataValidation(type="custom", formula1="LEN(H1)<5")
    for dv, ref in ((choice, "F1:F10"), (whole, "G1:G10"), (custom, "H1:H10")):
        ws.add_data_validation(dv)
        dv.add(ref)
    save(wb, "05-rules")


def notes():
    wb = Workbook()
    ws = wb.active
    ws.title = "메모"
    ws["A1"] = "메모 있음"
    ws["A1"].comment = Comment("한국어 메모입니다\n둘째 줄", "검토자")
    ws["B2"] = "누리집"
    ws["B2"].hyperlink = "https://example.com/경로?q=1"
    ws["B3"] = "안쪽 이동"
    ws["B3"].hyperlink = "#'메모'!A1"
    ws["B4"] = "편지"
    ws["B4"].hyperlink = "mailto:someone@example.com"
    save(wb, "06-notes-links")


def tables():
    wb = Workbook()
    ws = wb.active
    ws.title = "표"
    ws.append(["품목", "수량", "단가", "금액"])
    for i, name in enumerate(["사과", "배", "감", "귤"], 1):
        ws.append([name, i, 1000 * i, f"=B{i + 1}*C{i + 1}"])
    table = Table(displayName="판매표", ref="A1:D5")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(table)
    other = wb.create_sheet("필터")
    other.append(["이름", "점수"])
    for i in range(1, 8):
        other.append([f"학생{i}", i * 11])
    other.auto_filter.ref = "A1:B8"
    other.auto_filter.add_filter_column(1, ["22", "33"])
    save(wb, "07-tables-filter")


def charts():
    wb = Workbook()
    ws = wb.active
    ws.title = "차트"
    ws.append(["월", "매출", "비용"])
    for m, a, b in [("1월", 10, 7), ("2월", 14, 8), ("3월", 9, 9), ("4월", 17, 10)]:
        ws.append([m, a, b])
    data = Reference(ws, min_col=2, min_row=1, max_col=3, max_row=5)
    cats = Reference(ws, min_col=1, min_row=2, max_row=5)
    bar = BarChart()
    bar.title = "월별 매출"
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    ws.add_chart(bar, "E2")
    line = LineChart()
    line.title = "추이"
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    ws.add_chart(line, "E20")
    pie = PieChart()
    pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    pie.set_categories(cats)
    ws.add_chart(pie, "N2")
    save(wb, "08-charts")


def printing():
    wb = Workbook()
    ws = wb.active
    ws.title = "인쇄"
    for r in range(1, 60):
        ws.cell(r, 1, f"행 {r}")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.3
    ws.page_margins.top = 1.1
    ws.print_options.gridLines = True
    ws.print_options.horizontalCentered = True
    ws.oddHeader.center.text = "머리글 &P"
    ws.oddFooter.right.text = "바닥글"
    ws.print_area = "A1:A50"
    ws.print_title_rows = "1:2"
    ws.row_breaks.append(__import__("openpyxl").worksheet.pagebreak.Break(id=25))
    locked = wb.create_sheet("보호")
    locked["A1"] = "보호된 시트"
    locked.protection.sheet = True
    locked.protection.password = "1234"
    locked.protection.formatCells = False
    wb.security = __import__("openpyxl").workbook.protection.WorkbookProtection(lockStructure=True)
    save(wb, "09-print-protect")


for make in (values, formulas, styles, layout, rules, notes, tables, charts, printing):
    make()
print("written to", out)
